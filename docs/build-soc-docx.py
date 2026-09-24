#!/usr/bin/env python3
"""Build a Word document containing only the Service-Operations-Collaborators table."""
import re, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = "/Users/patiphonpuntusin/Time/Chula/sw-arch/hire-assist"
SRC = os.path.join(REPO, "docs", "SERVICE-OPERATIONS-COLLABORATORS.md")
OUT = os.path.join(REPO, "docs", "HireAssist-SOC.docx")

INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1B, 0x5E, 0x8C)
MUTED = RGBColor(0x55, 0x5F, 0x6B)

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]*\)|\*[^*\n]+?\*)')


def shade(el, fill):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), fill)
    el.append(sh)


def add_runs(par, text, italic=False, bold=False, color=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            add_runs(par, tok[2:-2], italic=italic, bold=True, color=color); continue
        if tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            add_runs(par, tok[1:-1], italic=True, bold=bold, color=color); continue
        if tok.startswith('[') and '](' in tok:
            add_runs(par, tok[1:tok.index('](')], italic=italic, bold=bold, color=color); continue
        r = par.add_run(tok)
        r.font.size = Pt(8.5)
        if tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r.text = tok[1:-1]
            r.font.name = 'Consolas'
            r.font.size = Pt(8)
            r.font.color.rgb = ACCENT
        elif color is not None:
            r.font.color.rgb = color
        r.italic = italic
        r.bold = bold


def balance_across_breaks(md):
    """A *\u2026* or **\u2026** span may straddle a <br>. Close and reopen it on each line so
    that splitting on <br> never leaves an unpaired marker."""
    def fix(m, mark):
        parts = m.group(1).split('<br>')
        return '<br>'.join(f'{mark}{p}{mark}' if p.strip() else p for p in parts)
    md = re.sub(r'\*\*(.+?)\*\*', lambda m: fix(m, '**'), md, flags=re.S)
    md = re.sub(r'(?<!\*)\*([^*]+?)\*(?!\*)', lambda m: fix(m, '*'), md, flags=re.S)
    return md


def cell_text(cell, md):
    """Render one markdown table cell, honouring <br> as a line break."""
    md = balance_across_breaks(md.replace('\u3000', '   '))
    lines = md.split('<br>')
    cell.text = ''
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        if line.strip():
            add_runs(p, line.strip())
        else:
            p.add_run('').font.size = Pt(4)


def main():
    text = open(SRC).read()

    rows = []
    for line in text.split('\n'):
        s = line.strip()
        if s.startswith('|') and s.endswith('|'):
            if re.match(r'^\|[\s:\-\|]+\|$', s):
                continue
            rows.append([c.strip() for c in s.strip('|').split('|')])
        elif rows:
            break                      # first table only
    assert rows and len(rows[0]) == 3, "SOC table not found or shape changed"

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Inches(0.6)
    sec.top_margin = sec.bottom_margin = Inches(0.6)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(9)
    normal.font.color.rgb = INK

    h = doc.add_paragraph()
    h.paragraph_format.space_after = Pt(2)
    r = h.add_run('Service–Operations–Collaborators')
    r.font.size = Pt(17); r.bold = True; r.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    r = sub.add_run('HireAssist · five services behind one API gateway · every collaboration is '
                    'REST over HTTP/JSON')
    r.font.size = Pt(9.5); r.font.color.rgb = MUTED

    t = doc.add_table(rows=len(rows), cols=3)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = (Inches(2.5), Inches(3.0), Inches(4.3))

    for i, row in enumerate(rows):
        for j in range(3):
            c = t.cell(i, j)
            c.width = widths[j]
            if i == 0:
                c.text = ''
                p = c.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                rr = p.add_run(row[j])
                rr.bold = True; rr.font.size = Pt(9.5)
                shade(c._tc.get_or_add_tcPr(), 'E8EEF4')
            else:
                cell_text(c, row[j])
                tcPr = c._tc.get_or_add_tcPr()
                mar = OxmlElement('w:tcMar')
                for side, v in (('top', 90), ('bottom', 90), ('left', 110), ('right', 110)):
                    el = OxmlElement('w:' + side)
                    el.set(qn('w:w'), str(v)); el.set(qn('w:type'), 'dxa')
                    mar.append(el)
                tcPr.append(mar)

    # repeat the header row across pages
    trPr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true'); trPr.append(hdr)

    cp = doc.core_properties
    cp.title = 'HireAssist — Service–Operations–Collaborators'
    cp.author = 'Patiphon Puntusin, Thanabul Parodom, Thanwarat Korcharoenkiat, Rerngrit Jangsri'

    doc.save(OUT)
    print('wrote', OUT)
    print('rows:', len(rows) - 1, 'services')


if __name__ == '__main__':
    main()
